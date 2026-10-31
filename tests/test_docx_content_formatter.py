"""
Tests for PagePilot DOCX Content Formatter (Stage 5).
"""

import pytest
import docx
from pagepilot.docx.content_formatter import ContentFormatter, append_formatted_runs
from pagepilot.docx.templates import SIMPLE_STUDY_NOTES


def test_append_formatted_runs():
    doc = docx.Document()
    p = doc.add_paragraph()
    text = "Hello **bold** world *italic* and `code` test."
    append_formatted_runs(
        p,
        text,
        font_name="Calibri",
        font_size_pt=11.0,
        color=SIMPLE_STUDY_NOTES.color_text,
    )

    runs = p.runs
    assert len(runs) >= 4
    # Check bold run
    bold_runs = [r for r in runs if r.bold and r.text == "bold"]
    assert len(bold_runs) == 1

    # Check italic run
    italic_runs = [r for r in runs if r.italic and r.text == "italic"]
    assert len(italic_runs) == 1

    # Check code run
    code_runs = [r for r in runs if r.font.name == "Consolas" and r.text == "code"]
    assert len(code_runs) == 1


def test_render_headings_and_paragraphs():
    doc = docx.Document()
    formatter = ContentFormatter(doc, SIMPLE_STUDY_NOTES)

    md = """# Main Title
## Section 1
Standard paragraph text.
### Subsection
Another paragraph with **emphasis**.
"""
    formatter.render_markdown(md)

    headings = [p for p in doc.paragraphs if p.runs and p.runs[0].bold]
    assert len(headings) >= 3


def test_render_lists():
    doc = docx.Document()
    formatter = ContentFormatter(doc, SIMPLE_STUDY_NOTES)

    md = """- Bullet Item 1
- Bullet Item 2
1. Numbered Item 1
2. Numbered Item 2
"""
    formatter.render_markdown(md)

    bullet_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    num_paras = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(bullet_paras) == 2
    assert len(num_paras) == 2


def test_render_callout_box():
    doc = docx.Document()
    formatter = ContentFormatter(doc, SIMPLE_STUDY_NOTES)

    md = """> Note: This is an important callout definition.
"""
    formatter.render_markdown(md)

    # Callout should be rendered as a single-cell table
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 1
    assert len(table.columns) == 1
    assert "important callout definition" in table.rows[0].cells[0].text


def test_render_code_block():
    doc = docx.Document()
    formatter = ContentFormatter(doc, SIMPLE_STUDY_NOTES)

    md = """```
def sample_code():
    return 42
```"""
    formatter.render_markdown(md)

    assert len(doc.tables) == 1
    cell_text = doc.tables[0].rows[0].cells[0].text
    assert "def sample_code():" in cell_text


def test_render_markdown_table():
    doc = docx.Document()
    formatter = ContentFormatter(doc, SIMPLE_STUDY_NOTES)

    md = """| Col A | Col B |
| --- | --- |
| Val 1 | Val 2 |
| Val 3 | Val 4 |
"""
    formatter.render_markdown(md)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.rows[0].cells[0].text == "Col A"
    assert table.rows[1].cells[1].text == "Val 2"

