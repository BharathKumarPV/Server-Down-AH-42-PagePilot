"""
End-to-end Integration Tests for Stage 6A Multi-Document Reasoning and Verification.
"""

from pathlib import Path
from unittest.mock import MagicMock
import pytest
import docx

from pagepilot.ai.claim_models import (
    Claim,
    Contradiction,
    Evidence,
    VerificationStatus,
    VerifiedAnswer,
)
from pagepilot.ai.rag_service import RAGService
from pagepilot.docx.document_service import DocxService
from pagepilot.docx.validators import validate_docx_file
from pagepilot.evaluation.dataset import EVALUATION_DATASET
from pagepilot.pdf.models import Chunk


def test_evaluation_dataset_structure():
    assert len(EVALUATION_DATASET) == 14
    for case in EVALUATION_DATASET:
        assert case.case_id > 0
        assert len(case.question) > 5
        assert len(case.synthetic_docs) > 0


def test_docx_export_of_verified_multi_doc_answer(tmp_path):
    ev1 = Evidence(
        evidence_id="E1",
        document_id=1,
        filename="Security_Arch.pdf",
        page_number=3,
        chunk_id="c1",
        chunk_index=0,
        original_text="Symmetric encryption uses AES-256 for all stored datasets.",
    )
    ev2 = Evidence(
        evidence_id="E2",
        document_id=2,
        filename="Network_Arch.pdf",
        page_number=7,
        chunk_id="c2",
        chunk_index=0,
        original_text="Public-key cryptography uses RSA-4096 for authentication.",
    )

    claims = [
        Claim(
            claim_id="C1",
            text="Symmetric encryption uses AES-256 for all stored datasets",
            evidence_ids=["E1"],
            verification_status=VerificationStatus.SUPPORTED,
        ),
        Claim(
            claim_id="C2",
            text="Public-key cryptography uses RSA-4096 for authentication",
            evidence_ids=["E2"],
            verification_status=VerificationStatus.SUPPORTED,
        ),
    ]

    contradictions = [
        Contradiction(
            topic="Worker Threads",
            statement_a="Worker threads is 8",
            source_a="[Security_Arch.pdf, p. 3]",
            doc_id_a=1,
            page_a=3,
            statement_b="Worker threads is 16",
            source_b="[Network_Arch.pdf, p. 7]",
            doc_id_b=2,
            page_b=7,
            explanation="Different concurrency settings across environments.",
        )
    ]

    ans = VerifiedAnswer(
        question="Compare encryption and concurrency across architectures.",
        claims=claims,
        evidence_map={"E1": ev1, "E2": ev2},
        contradictions=contradictions,
        unanswered_requirements=["Legacy protocols"],
        contributing_documents=["Security_Arch.pdf", "Network_Arch.pdf"],
    )

    docx_service = DocxService()
    out_file = tmp_path / "multi_doc_verified_report.docx"

    report_path = docx_service.generate_multi_document_report(
        verified_answer=ans,
        template_name="Project Report",
        output_path=out_file,
    )

    assert report_path.exists()
    assert validate_docx_file(report_path) is True

    doc = docx.Document(str(report_path))
    doc_text = " ".join(p.text for p in doc.paragraphs)

    # Check that sentences have their source citations
    assert "Security_Arch.pdf, p. 3" in doc_text
    assert "Network_Arch.pdf, p. 7" in doc_text
    # Check contradictions section
    assert "Cross-Document Contradictions" in doc_text
    # Check sources table
    assert len(doc.tables) >= 1


def test_answer_question_verified_pipeline_mocked_ollama():
    mock_client = MagicMock()
    mock_retriever = MagicMock()
    mock_db = MagicMock()

    chunk1 = Chunk("c1", 1, "DocA.pdf", 2, 0, "Symmetric encryption uses a single shared key.", 10)
    chunk2 = Chunk("c2", 2, "DocB.pdf", 4, 0, "Asymmetric encryption uses public and private key pairs.", 10)

    # Mock cross-document retrieval
    ev1 = Evidence("E1", 1, "DocA.pdf", 2, "c1", 0, chunk1.text)
    ev2 = Evidence("E2", 2, "DocB.pdf", 4, "c2", 0, chunk2.text)
    from pagepilot.ai.claim_models import InformationRequirement
    req1 = InformationRequirement("R1", "symmetric", has_sufficient_evidence=True)
    req2 = InformationRequirement("R2", "asymmetric", has_sufficient_evidence=True)

    mock_retriever.retrieve_cross_document.return_value = ([ev1, ev2], [req1, req2])

    # Mock Ollama JSON streaming output
    mock_json = """{
        "claims": [
            {"claim_id": "C1", "text": "Symmetric encryption uses a single shared key.", "claim_type": "factual", "evidence_ids": ["E1"]},
            {"claim_id": "C2", "text": "Asymmetric encryption uses public and private key pairs.", "claim_type": "factual", "evidence_ids": ["E2"]}
        ],
        "unanswered_requirements": []
    }"""
    mock_client.chat_stream.return_value = iter([mock_json])

    service = RAGService(ollama_client=mock_client, hybrid_retriever=mock_retriever, db=mock_db)

    verified_ans = service.answer_question_verified(
        question="Compare symmetric and asymmetric encryption.",
        document_ids=[1, 2],
    )

    assert len(verified_ans.claims) == 2
    assert all(c.is_verified for c in verified_ans.claims)
    assert set(verified_ans.contributing_documents) == {"DocA.pdf", "DocB.pdf"}

    rendered = verified_ans.render_sentence_level_text()
    assert "[DocA.pdf, p. 2]" in rendered
    assert "[DocB.pdf, p. 4]" in rendered

