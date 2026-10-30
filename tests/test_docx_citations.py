"""
Tests for PagePilot DOCX citations and reference table rendering (Stage 5).
"""

import pytest
import docx
from pagepilot.docx.citation_formatter import (
    CitationMode,
    add_sources_section,
    format_inline_citation,
)
from pagepilot.docx.templates import ACADEMIC_SUMMARY, SIMPLE_STUDY_NOTES


def test_format_inline_citation():
    assert format_inline_citation("Report.pdf", 3) == "[Report.pdf, p. 3]"
    assert format_inline_citation("Handbook.pdf", None) == "[Handbook.pdf]"
    assert format_inline_citation("Manual.pdf", 0) == "[Manual.pdf]"
    assert format_inline_citation("", 1) == "[Document.pdf, p. 1]"


def test_citation_mode_logic():
    assert CitationMode.should_include_inline(CitationMode.BOTH) is True
    assert CitationMode.should_include_inline(CitationMode.INLINE_ONLY) is True
    assert CitationMode.should_include_inline(CitationMode.SOURCES_ONLY) is False
    assert CitationMode.should_include_inline(CitationMode.NONE) is False

    assert CitationMode.should_include_sources_section(CitationMode.BOTH) is True
    assert CitationMode.should_include_sources_section(CitationMode.SOURCES_ONLY) is True
    assert CitationMode.should_include_sources_section(CitationMode.INLINE_ONLY) is False
    assert CitationMode.should_include_sources_section(CitationMode.NONE) is False


def test_add_sources_section():
    doc = docx.Document()
    citations = [
        {"filename": "DocA.pdf", "page_number": 2, "snippet": "First important passage."},
        {"filename": "DocA.pdf", "page_number": 2, "snippet": "Duplicate entry to deduplicate."},
        {"filename": "DocB.pdf", "page_number": 5, "snippet": "Second distinct reference."},
    ]

    add_sources_section(doc, citations, ACADEMIC_SUMMARY)

    # Verify table added
    assert len(doc.tables) == 1
    table = doc.tables[0]
    # Header row + 2 unique rows (since duplicate page 2 in DocA is deduplicated)
    assert len(table.rows) == 3

    # Check headers
    hdr_cells = table.rows[0].cells
    assert hdr_cells[0].text == "#"
    assert hdr_cells[1].text == "Source Document"
    assert hdr_cells[2].text == "Page"
    assert "Reference" in hdr_cells[3].text

    # Check row 1
    r1 = table.rows[1].cells
    assert r1[0].text == "1"
    assert r1[1].text == "DocA.pdf"
    assert r1[2].text == "p. 2"

    # Check row 2
    r2 = table.rows[2].cells
    assert r2[0].text == "2"
    assert r2[1].text == "DocB.pdf"
    assert r2[2].text == "p. 5"


def test_add_sources_section_empty():
    doc = docx.Document()
    add_sources_section(doc, [], SIMPLE_STUDY_NOTES)
    assert len(doc.tables) == 0

