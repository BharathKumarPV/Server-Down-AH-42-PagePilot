"""
Tests for PagePilot DocxService (Stage 5).
"""

from pathlib import Path
import pytest
import docx

from pagepilot.db.database import DocumentDB
from pagepilot.docx.document_service import DocxService
from pagepilot.docx.citation_formatter import CitationMode
from pagepilot.pdf.models import Chunk


@pytest.fixture
def test_db(tmp_path):
    db_file = tmp_path / "test_docs.db"
    db = DocumentDB(db_path=db_file)
    doc = db.create_document(
        filename="Security_Report.pdf",
        file_path=str(tmp_path / "Security_Report.pdf"),
        sha256_hash="abc123hash",
        page_count=5,
        embed_model="nomic-embed-text",
    )
    doc_id = doc.id
    chunks_data = [
        Chunk(
            chunk_id=f"sec_c{i}",
            document_id=doc_id,
            filename="Security_Report.pdf",
            page_number=i,
            chunk_index=i - 1,
            text=f"Content for security chunk number {i} with vital details on protocol and analysis.",
            token_count=15,
            section_heading=f"Security Section {i}",
        )
        for i in range(1, 6)
    ]
    db.insert_chunks(chunks_data)
    return db, doc_id


def test_service_generate_pdf_summary(test_db, tmp_path):
    db, doc_id = test_db
    service = DocxService(db=db)
    out_file = tmp_path / "summary_test.docx"

    res = service.generate_pdf_summary(
        doc_id=doc_id,
        template_name="Academic Summary",
        output_path=out_file,
    )

    assert res.exists()
    doc = docx.Document(str(res))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Executive Summary" in text
    assert "Security_Report.pdf" in text


def test_service_generate_study_notes(test_db, tmp_path):
    db, doc_id = test_db
    service = DocxService(db=db)
    out_file = tmp_path / "notes_test.docx"

    res = service.generate_study_notes(
        doc_id=doc_id,
        topic="Network Protocols",
        template_name="Simple Study Notes",
        output_path=out_file,
    )

    assert res.exists()
    doc = docx.Document(str(res))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Study Notes" in text
    assert "Network Protocols" in text


def test_service_generate_qa_document(tmp_path):
    service = DocxService()
    out_file = tmp_path / "qa_test.docx"

    qa_pairs = [
        {
            "question": "What is the primary encryption standard?",
            "answer": "AES-256 is the mandated standard for data at rest.",
            "citations": [{"filename": "Standard.pdf", "page_number": 4, "snippet": "AES-256 standard."}],
        },
        {
            "question": "How are session tokens refreshed?",
            "answer": "Short-lived JWTs are renewed with cryptographic rotation.",
            "citations": [{"filename": "Standard.pdf", "page_number": 7, "snippet": "Cryptographic rotation."}],
        },
    ]

    res = service.generate_qa_document(
        qa_pairs=qa_pairs,
        template_name="Question and Answer Notes",
        output_path=out_file,
    )

    assert res.exists()
    doc = docx.Document(str(res))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "What is the primary encryption standard?" in text
    assert "AES-256 is the mandated standard" in text


def test_service_generate_custom_document(tmp_path):
    service = DocxService()
    out_file = tmp_path / "custom_test.docx"

    res = service.generate_custom_document(
        title="Custom Operations Manual",
        content="# Section 1\nCustom workflow steps.",
        template_name="Custom Document",
        output_path=out_file,
    )

    assert res.exists()
    doc = docx.Document(str(res))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Custom Operations Manual" in text


def test_service_generate_selected_pages(test_db, tmp_path):
    db, doc_id = test_db
    service = DocxService(db=db)
    out_file = tmp_path / "selected_test.docx"

    res = service.generate_selected_pages(
        doc_id=doc_id,
        page_numbers=[2, 4],
        template_name="Simple Study Notes",
        output_path=out_file,
    )

    assert res.exists()
    doc = docx.Document(str(res))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Page 2" in text
    assert "Page 4" in text


def test_service_errors(test_db):
    db, doc_id = test_db
    service = DocxService(db=db)

    # Invalid doc id
    with pytest.raises(ValueError, match="not found in database"):
        service.generate_pdf_summary(doc_id=999999)

    # Invalid page numbers
    with pytest.raises(ValueError, match="No indexed content found for pages"):
        service.generate_selected_pages(doc_id=doc_id, page_numbers=[99])
