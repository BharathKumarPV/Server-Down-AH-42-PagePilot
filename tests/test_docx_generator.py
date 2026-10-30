"""
Tests for PagePilot DocxGenerator core engine (Stage 5).
"""

from pathlib import Path
import pytest
import docx

from pagepilot.docx.generator import DocxGenerator
from pagepilot.docx.citation_formatter import CitationMode
from pagepilot.docx.templates import TEMPLATE_NAMES


def test_generator_all_templates(tmp_path):
    generator = DocxGenerator()
    content = """# Executive Summary
This is a paragraph with **bold** text and *italic* text.

## Key Concepts
- Concept A: Description
- Concept B: Another description

> Important callout block with critical guidance.

| Header 1 | Header 2 |
| --- | --- |
| Data 1 | Data 2 |
"""
    citations = [
        {"filename": "Policy.pdf", "page_number": 3, "snippet": "Core policy excerpt."}
    ]

    for tmpl_name in TEMPLATE_NAMES:
        out_file = tmp_path / f"test_{tmpl_name.replace(' ', '_')}.docx"
        result_path = generator.generate(
            title=f"Test Document - {tmpl_name}",
            content=content,
            subtitle="Automated Verification Run",
            template_name=tmpl_name,
            citations=citations,
            citation_mode=CitationMode.BOTH,
            output_path=out_file,
        )

        assert result_path.exists()
        assert result_path.stat().st_size > 0

        # Verify reopening with docx
        reopened = docx.Document(str(result_path))
        assert len(reopened.paragraphs) > 0
        # Check title exists in document
        title_found = any(f"Test Document - {tmpl_name}" in p.text for p in reopened.paragraphs)
        assert title_found is True


def test_generator_citations_modes(tmp_path):
    generator = DocxGenerator()
    citations = [{"filename": "Guide.pdf", "page_number": 1, "snippet": "Guidance excerpt."}]

    # 1. Mode: SOURCES_ONLY (sources table present)
    p1 = generator.generate(
        title="Sources Only Doc",
        content="Plain content.",
        citations=citations,
        citation_mode=CitationMode.SOURCES_ONLY,
        output_path=tmp_path / "sources_only.docx",
    )
    doc1 = docx.Document(str(p1))
    assert len(doc1.tables) >= 1  # Sources table present

    # 2. Mode: NONE (no sources section)
    p2 = generator.generate(
        title="None Doc",
        content="Plain content without sources table.",
        citations=citations,
        citation_mode=CitationMode.NONE,
        output_path=tmp_path / "none_mode.docx",
    )
    doc2 = docx.Document(str(p2))
    assert len(doc2.tables) == 0


def test_generator_special_unicode_characters(tmp_path):
    generator = DocxGenerator()
    content = """# Special Unicode: ©, ®, ™, €, £, ¥, •, —, α, β, γ
Mathematical notations: E = mc², H₂O, ∑, ∫, ≈, ≠, ≤, ≥.
Languages: Привет, こんにちは, 你好, مرحبا.
"""
    out_file = tmp_path / "unicode_test.docx"
    result = generator.generate(
        title="Unicode & Symbols Test",
        content=content,
        output_path=out_file,
    )
    assert result.exists()
    doc = docx.Document(str(result))
    text_content = " ".join(p.text for p in doc.paragraphs)
    assert "©" in text_content
    assert "E = mc²" in text_content
    assert "こんにちは" in text_content

